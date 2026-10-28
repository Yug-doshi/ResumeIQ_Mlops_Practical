"""Run a synthetic ranking benchmark against the running ResumeIQ API."""

import os
import time
from io import BytesIO

import mlflow
import requests
from docx import Document

from mlops.evaluation import evaluate_ranking


API_URL = os.getenv("RESUMEIQ_API_URL", "http://localhost:8000")
TRACKING_URI = os.getenv("MLFLOW_TRACKING_URI", "http://localhost:5000")
EXPERIMENT_NAME = "ResumeIQ-Resume-Screening"
K = 3

# Synthetic benchmark only. Grades are assigned before seeing model results:
# 0 = not relevant, 1 = partially relevant, 2 = highly relevant.
JOB = {
    "job_title": "Python Backend Engineer",
    "job_description": (
        "Develop Python backend APIs using FastAPI, Docker, and PostgreSQL. "
        "Build reliable services, write automated tests, and deploy applications."
    ),
    "required_skills": "Python, FastAPI, Docker",
    "preferred_skills": "PostgreSQL, pytest",
}

BENCHMARK = [
    {
        "filename": "candidate_a.docx",
        "grade": 2,
        "text": (
            "Backend engineer experienced in Python, FastAPI, Docker, "
            "PostgreSQL, pytest, REST APIs, automated testing and deployment."
        ),
    },
    {
        "filename": "candidate_b.docx",
        "grade": 0,
        "text": (
            "Graphic designer experienced in illustration, typography, "
            "branding, visual design and photo editing."
        ),
    },
    {
        "filename": "candidate_c.docx",
        "grade": 1,
        "text": (
            "Software developer experienced in Python scripting, SQL, "
            "data processing and basic automated testing."
        ),
    },
    {
        "filename": "candidate_d.docx",
        "grade": 2,
        "text": (
            "Python API developer using FastAPI, Docker, PostgreSQL, "
            "pytest, REST services, CI pipelines and cloud deployment."
        ),
    },
]


def make_docx(text):
    document = Document()
    document.add_heading("Synthetic Resume", level=1)
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def main():
    files = [
        (
            "resumes",
            (item["filename"], make_docx(item["text"]),
             "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        )
        for item in BENCHMARK
    ]

    response = requests.post(
        f"{API_URL}/api/v1/match",
        data=JOB,
        files=files,
        timeout=300,
    )
    response.raise_for_status()
    result = response.json()

    if result["summary"]["processed"] != len(BENCHMARK):
        raise RuntimeError(
            f"Expected {len(BENCHMARK)} processed resumes; "
            f"got {result['summary']['processed']}. Errors: {result['errors']}"
        )

    grades = {item["filename"]: item["grade"] for item in BENCHMARK}
    ranked_ids = [candidate["filename"] for candidate in result["candidates"]]
    metrics = evaluate_ranking(ranked_ids, grades, k=K)

    mlflow.set_tracking_uri(TRACKING_URI)
    mlflow.set_experiment(EXPERIMENT_NAME)

    with mlflow.start_run(run_name="synthetic-ranking-evaluation"):
        mlflow.set_tags({
            "project": "ResumeIQ",
            "task": "ranking_evaluation",
            "dataset_type": "synthetic",
            "evaluation_kind": "smoke_benchmark",
        })
        mlflow.log_params({
            "model_id": result["model_id"],
            "scoring_version": result["scoring_version"],
            "skill_dictionary_version": result["skill_dictionary_version"],
            "candidate_count": len(BENCHMARK),
            "k": K,
            "job_profile": "python_backend_engineer",
        })
        mlflow.log_metrics(metrics)

    print("Model:", result["model_id"])
    print("Scoring version:", result["scoring_version"])
    print("\nRanked candidates:")
    for candidate in result["candidates"]:
        print(
            f"Rank {candidate['rank']}: {candidate['filename']} "
            f"(score={candidate['score']:.2f}, "
            f"reference_grade={grades[candidate['filename']]})"
        )

    print("\nRanking metrics:")
    for name, value in metrics.items():
        print(f"{name}: {value:.4f}")

    print("\nEvaluation run logged to MLflow:", EXPERIMENT_NAME)
    print("Note: synthetic benchmark results are not real-world accuracy estimates.")


if __name__ == "__main__":
    main()
